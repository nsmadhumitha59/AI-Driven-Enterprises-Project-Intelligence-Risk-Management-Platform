"""
Generates rich enterprise sample project documents for PDF, DOCX, CSV, and TXT testing & demonstration.
"""
import os
from pathlib import Path
import pandas as pd
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

OUTPUT_DIR = Path(__file__).resolve().parent.parent / "sample_documents"


def generate_all_samples():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    print(f"Creating sample documents in: {OUTPUT_DIR}")

    # 1. DOCX: Project Charter & PRD
    docx_path = OUTPUT_DIR / "Project_Charter_PRD.docx"
    doc = Document()
    doc.add_heading("Enterprise Cloud Analytics Platform — Project Charter & PRD", level=0)
    
    doc.add_heading("1. Executive Summary & Project Goals", level=1)
    doc.add_paragraph(
        "The Enterprise Cloud Analytics Platform is an autonomous, AI-driven project management intelligence system. "
        "The project goal is to reduce milestone slippage by 40% and provide continuous risk forecasting across "
        "engineering deliverables, sprint backlogs, and multi-format document artifacts."
    )
    doc.add_paragraph(
        "Key Project Objectives:\n"
        "- GOAL-101: Autonomous multi-agent extraction of scope, timeline risks, and active blockers.\n"
        "- GOAL-102: Sub-second semantic vector retrieval grounded in persistent ChromaDB vector storage.\n"
        "- GOAL-103: Automated Agile User Story generation and multi-dimensional project health scorecard diagnostics.\n"
        "- GOAL-104: Incremental document ingestion with zero data loss or index degradation."
    )

    doc.add_heading("2. Scope Boundaries & Technical Constraints", level=1)
    doc.add_paragraph(
        "In Scope:\n"
        "- Ingestion of PDF, DOCX, CSV, and TXT/Markdown documents.\n"
        "- SentenceTransformers semantic vector indexing with cosine similarity retrieval.\n"
        "- Multi-agent intelligence pipeline (Scope, Risk, Blocker, Documentation, Health Scorer, Chat).\n"
        "- Interactive Executive Dashboard with Risk Priority Matrix scatter plot.\n\n"
        "Out of Scope:\n"
        "- Direct legacy Jira SOAP integration (replaced by standardized CSV exports).\n"
        "- Hardware on-premise appliance deployments (Cloud/Local containerized runtime only)."
    )

    doc.add_heading("3. Key Milestones & Responsibility Matrix", level=1)
    table = doc.add_table(rows=1, cols=4)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Milestone / Deliverable"
    hdr_cells[1].text = "Owner"
    hdr_cells[2].text = "Target Date"
    hdr_cells[3].text = "Status"

    milestones_data = [
        ("M1: Ingestion & Vector RAG Subsystem", "Sarah Connor (Data Lead)", "2026-10-15", "Completed"),
        ("M2: Multi-Agent Intelligence Network", "Marcus Vance (AI Lead)", "2026-10-25", "In Progress"),
        ("M3: Health Scoring & Conversational Q&A", "Elena Rostova (Platform Eng)", "2026-11-05", "In Progress"),
        ("M4: Executive Dashboard & Incremental Sync", "David Chen (Frontend Lead)", "2026-11-20", "Active"),
    ]

    for item in milestones_data:
        row_cells = table.add_row().cells
        row_cells[0].text = item[0]
        row_cells[1].text = item[1]
        row_cells[2].text = item[2]
        row_cells[3].text = item[3]

    doc.save(str(docx_path))
    print(f"✅ Generated DOCX: {docx_path.name}")

    # 2. PDF: System Architecture & Risk Specification
    pdf_path = OUTPUT_DIR / "System_Architecture_Spec.pdf"
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)

    pdf_stream_text = (
        b"BT\n"
        b"/Helvetica-Bold 14 Tf\n"
        b"50 740 Td\n"
        b"(Enterprise Architecture Decision Record & Risk Assessment) Tj\n"
        b"/Helvetica 11 Tf\n"
        b"0 -24 Td\n"
        b"(System Architecture Overview:) Tj\n"
        b"0 -16 Td\n"
        b"(1. RAG Vector Subsystem: SentenceTransformers all-MiniLM-L6-v2 with ChromaDB persistent storage.) Tj\n"
        b"0 -16 Td\n"
        b"(2. Agent Orchestration: Multi-agent pipeline coordinating Scope, Risk, and Blocker agents.) Tj\n"
        b"0 -16 Td\n"
        b"(3. API Gateway: FastAPI backend with CORS, OpenAPI schemas, and async background tasks.) Tj\n"
        b"0 -24 Td\n"
        b"(Technical Risks & Mitigation Strategies:) Tj\n"
        b"0 -16 Td\n"
        b"(RSK-01 [CRITICAL]: Payment Gateway OAuth2 token timeout causes transaction failure.) Tj\n"
        b"0 -16 Td\n"
        b"(    Mitigation: Implement distributed Redis token caching with automated renewal.) Tj\n"
        b"0 -16 Td\n"
        b"(RSK-02 [HIGH]: ChromaDB sqlite locking on high concurrent batch ingestion.) Tj\n"
        b"0 -16 Td\n"
        b"(    Mitigation: Wrap vector mutations in serialized thread locks with retry backoff.) Tj\n"
        b"0 -16 Td\n"
        b"(RSK-03 [MEDIUM]: Large PDF parsing memory footprint exceeding container thresholds.) Tj\n"
        b"0 -16 Td\n"
        b"(    Mitigation: Stream PDF page extraction lazily via pypdf chunk iterators.) Tj\n"
        b"0 -16 Td\n"
        b"(RSK-04 [LOW]: Static asset caching headers causing delayed CSS dashboard reloads.) Tj\n"
        b"0 -16 Td\n"
        b"(    Mitigation: Configure Cache-Control max-age=0 for dynamic portal development.) Tj\n"
        b"ET\n"
    )

    stream_obj = DecodedStreamObject()
    stream_obj.set_data(pdf_stream_text)
    page[NameObject("/Contents")] = writer._add_object(stream_obj)

    # Add fonts
    font_bold = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica-Bold"),
    })
    font_regular = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    fonts_res = DictionaryObject({
        NameObject("/Helvetica-Bold"): writer._add_object(font_bold),
        NameObject("/Helvetica"): writer._add_object(font_regular),
    })
    resources = DictionaryObject({
        NameObject("/Font"): writer._add_object(fonts_res)
    })
    page[NameObject("/Resources")] = writer._add_object(resources)

    with open(pdf_path, "wb") as f:
        writer.write(f)
    print(f"✅ Generated PDF: {pdf_path.name}")

    # 3. CSV: Sprint Backlog Tasks & Delivery Tracking
    csv_path = OUTPUT_DIR / "Sprint_Backlog_Tasks.csv"
    tasks_data = {
        "Task_ID": ["TASK-301", "TASK-302", "TASK-303", "TASK-304", "TASK-305", "TASK-306", "TASK-307"],
        "Title": [
            "Implement Recursive Boundary Text Splitter",
            "Integrate Stripe & PayPal Webhook Handlers",
            "Design Risk Priority Matrix SVG Scatter Plot",
            "Configure ChromaDB Vector Collection Persistence",
            "Implement Zero-Hallucination Grounded Chatbot",
            "Perform Load Testing on Batch Document Ingestion",
            "Build Multi-Dimensional Health Scoring Formula"
        ],
        "Category": ["Backend", "Integration", "Frontend", "Database", "AI/NLP", "QA", "Analytics"],
        "Assignee": ["Alice Smith", "Bob Vance", "David Chen", "Alice Smith", "Elena Rostova", "Marcus Vance", "Elena Rostova"],
        "Priority": ["High", "Critical", "High", "Critical", "High", "Medium", "High"],
        "Status": ["Completed", "Blocked", "Completed", "Completed", "In Progress", "In Progress", "Completed"],
        "Blocker_Reason": [
            "None",
            "Awaiting sandbox API keys from Third-Party Security Compliance team",
            "None",
            "None",
            "None",
            "Requires staging environment Kubernetes cluster provisioning",
            "None"
        ],
        "Due_Date": ["2026-10-18", "2026-10-22", "2026-10-24", "2026-10-19", "2026-10-28", "2026-10-30", "2026-10-20"]
    }
    df = pd.DataFrame(tasks_data)
    df.to_csv(csv_path, index=False)
    print(f"✅ Generated CSV: {csv_path.name}")

    # 4. TXT: Sprint Standup & Architecture Meeting Minutes
    txt_path = OUTPUT_DIR / "Sprint_Standup_Meeting_Notes.txt"
    notes_content = (
        "Project Intelligence Engineering — Sprint 4 Standup & Architecture Minutes\n"
        "Date: October 19, 2026 | Facilitator: Marcus Vance | Attendees: Alice, Bob, David, Elena\n\n"
        "1. Active Blockers & Critical Dependencies:\n"
        "- BLOCKER-01: Third-party payment gateway sandbox credentials remain blocked by InfoSec review (Assigned: Bob Vance, Resolution ETA: 2026-10-22).\n"
        "- BLOCKER-02: Staging cluster VPC peering latency spikes exceeding 250ms SLA (Assigned: Sarah Connor, Resolution ETA: 2026-10-23).\n\n"
        "2. Action Items & Next Steps:\n"
        "- ACTION-101: David Chen to finalize Risk Priority Matrix bubble hover tooltips and inspector panel by end of day.\n"
        "- ACTION-102: Elena Rostova to benchmark Health Scorer weights across 50 test project repositories.\n"
        "- ACTION-103: Alice Smith to add incremental ingestion integrity checks for duplicate document hashes.\n"
        "- ACTION-104: Marcus Vance to review prompt guardrails for conversational assistant to prevent out-of-domain answers.\n\n"
        "3. Delivery Trajectory Forecast:\n"
        "Current sprint velocity is 85% on target. Delivery trajectory is evaluated as 'At Risk' due to payment gateway blocker, "
        "with an estimated 25% schedule slippage risk unless sandbox keys are unblocked before Wednesday."
    )
    txt_path.write_text(notes_content, encoding="utf-8")
    print(f"✅ Generated TXT: {txt_path.name}")

    print("\n🎉 All 4 multi-format enterprise sample documents generated successfully!")


if __name__ == "__main__":
    generate_all_samples()
